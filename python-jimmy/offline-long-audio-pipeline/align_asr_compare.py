"""Align k2 vs pipeline segments by time interval and write a Chinese comparison workbook."""

from __future__ import annotations

from dataclasses import dataclass
import json
import re
from pathlib import Path
from typing import Any


@dataclass
class SegmentView:
    segment_id: str
    start_ms: int
    end_ms: int
    speaker: str
    language: str
    asr_text: str
    zh_text: str = ""


def _format_clock(milliseconds: int) -> str:
    minutes, remainder = divmod(max(0, milliseconds), 60_000)
    seconds, millis = divmod(remainder, 1_000)
    return f"{minutes}:{seconds:02d}.{millis:03d}"


def load_k2_segments(path: Path) -> list[SegmentView]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    segments = []
    for item in payload.get("segments", []):
        start_ms = int(item["offsetMs"])
        duration_ms = int(item["durationMs"])
        segments.append(
            SegmentView(
                segment_id=str(item.get("segmentId") or ""),
                start_ms=start_ms,
                end_ms=start_ms + duration_ms,
                speaker=str(item.get("speakerId") or ""),
                language=str(item.get("asrLanguage") or ""),
                asr_text=str(item.get("asrText") or "").strip(),
            )
        )
    return segments


def _parse_pipeline_bounds(item: dict[str, Any]) -> tuple[int, int]:
    segment_id = str(item.get("segment_id") or "")
    match = re.match(r"^\d+_(\d+)_(\d+)$", segment_id)
    if match:
        return int(match.group(1)), int(match.group(2))
    raise ValueError(f"Cannot parse pipeline segment_id: {segment_id}")


def load_pipeline_segments(path: Path) -> list[SegmentView]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    segments = []
    for item in payload.get("segments", []):
        start_ms, end_ms = _parse_pipeline_bounds(item)
        segments.append(
            SegmentView(
                segment_id=str(item.get("segment_id") or ""),
                start_ms=start_ms,
                end_ms=end_ms,
                speaker=str(item.get("speaker_id") or ""),
                language=str(item.get("asr_language") or ""),
                asr_text=str(item.get("asr_text") or "").strip(),
            )
        )
    return segments


def covering(segments: list[SegmentView], start_ms: int, end_ms: int) -> SegmentView | None:
    overlap = []
    for segment in segments:
        left = max(segment.start_ms, start_ms)
        right = min(segment.end_ms, end_ms)
        if right > left:
            overlap.append((right - left, segment))
    if not overlap:
        return None
    overlap.sort(key=lambda item: (-item[0], item[1].start_ms, item[1].segment_id))
    return overlap[0][1]


def align_segments(
    before: list[SegmentView], after: list[SegmentView]
) -> list[tuple[int, int, SegmentView | None, SegmentView | None]]:
    points = sorted({0, *(item.start_ms for item in before + after), *(item.end_ms for item in before + after)})
    rows: list[tuple[int, int, SegmentView | None, SegmentView | None]] = []
    for start_ms, end_ms in zip(points, points[1:]):
        left = covering(before, start_ms, end_ms)
        right = covering(after, start_ms, end_ms)
        if left is None and right is None:
            continue
        if (
            rows
            and rows[-1][2] is left
            and rows[-1][3] is right
            and rows[-1][1] == start_ms
        ):
            previous = rows[-1]
            rows[-1] = (previous[0], end_ms, left, right)
        else:
            rows.append((start_ms, end_ms, left, right))
    return rows


def _side_cells(segment: SegmentView | None, previous: SegmentView | None) -> tuple[str, str, str]:
    """Keep the first row of a 1-to-many group; later rows on that side stay blank."""
    if segment is None or segment is previous:
        return ("", "", "")
    return (segment.segment_id, segment.speaker, segment.zh_text)


def write_xlsx(
    path: Path,
    rows: list[tuple[int, int, SegmentView | None, SegmentView | None]],
) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
    from openpyxl.utils import get_column_letter

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "片段对照"
    sheet.merge_cells("A1:C1")
    sheet.merge_cells("D1:F1")
    sheet["A1"] = "对比前（老脚本 k2）"
    sheet["D1"] = "对比后（新脚本 pipeline）"
    sheet.append(["片段ID", "说话人ID", "中文翻译", "片段ID", "说话人ID", "中文翻译"])

    header_fill = PatternFill("solid", fgColor="17324D")
    header_font = Font(color="FFFFFF", bold=True)
    before_fill = PatternFill("solid", fgColor="F4F7F9")
    after_fill = PatternFill("solid", fgColor="FBFDFD")
    thin = Border(
        left=Side(style="thin", color="D7E0E8"),
        right=Side(style="thin", color="D7E0E8"),
        top=Side(style="thin", color="D7E0E8"),
        bottom=Side(style="thin", color="D7E0E8"),
    )
    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="center")
    for row_index in (1, 2):
        for cell in sheet[row_index]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center
            cell.border = thin

    previous_left = None
    previous_right = None
    for index, (_start_ms, _end_ms, left, right) in enumerate(rows, start=1):
        values = [
            *_side_cells(left, previous_left),
            *_side_cells(right, previous_right),
        ]
        sheet.append(values)
        excel_row = index + 2
        for column, cell in enumerate(sheet[excel_row], start=1):
            cell.border = thin
            cell.alignment = wrap
            cell.fill = before_fill if column <= 3 else after_fill
        previous_left = left
        previous_right = right

    widths = [26, 14, 56, 26, 14, 56]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:F{sheet.max_row}"
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)


def apply_translations(segments: list[SegmentView], mapping: dict[str, str]) -> None:
    for segment in segments:
        key = segment.asr_text.strip()
        if key in mapping:
            segment.zh_text = mapping[key]
        elif not key:
            segment.zh_text = ""
        else:
            segment.zh_text = mapping.get(key, "")
