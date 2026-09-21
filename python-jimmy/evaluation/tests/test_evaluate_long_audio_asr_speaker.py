from __future__ import annotations

import json
import sys
import tempfile
import unittest
from pathlib import Path

EVALUATION_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(EVALUATION_DIR))
import evaluate_long_audio_asr_speaker as runner  # noqa: E402


class ResultAndLabelParsingTest(unittest.TestCase):
    def _write_result(self, root: Path) -> Path:
        run = root / "run"
        run.mkdir()
        (run / "result.json").write_text(
            json.dumps(
                {
                    "audio_name": "23_asr_1782715267098.pcm",
                    "segments": [
                        {
                            "segment_id": "23_asr_1782715267098_1118_2308",
                            "duration_ms": 2308,
                            "speaker_id": "speaker_02",
                            "asr_text": "\u6765\u5427",
                        },
                        {
                            "segment_id": "23_asr_1782715267098_3426_6846",
                            "duration_ms": 6846,
                            "speaker_id": "speaker_00",
                            "asr_text": "\u53ef\u4ee5",
                        },
                    ],
                },
                ensure_ascii=False,
            ),
            encoding="utf-8",
        )
        return run

    def test_parse_result_resolve_label_and_write_public_projection(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            results_dir = self._write_result(root)
            labels_dir = root / "labels"
            labels_dir.mkdir()
            (labels_dir / "23_asr_1782715267098_label.txt").write_text(
                "23_asr_1782715267098_1118_2335 (leslie) \u6765\u5427\n"
                "23_asr_1782715267098_3453_6823 (hui) \u53ef\u4ee5\n",
                encoding="utf-8",
            )
            records = runner.discover_result_records(results_dir)
            self.assertEqual(records[0].file_id, "23_asr_1782715267098")
            self.assertEqual(records[0].segments[0].start_ms, 1118)
            self.assertEqual(records[0].segments[0].end_ms, 3426)
            labels = runner.resolve_labels(records, labels_dir, explicit_label=None)
            self.assertEqual(
                labels[records[0].file_id].path.name,
                "23_asr_1782715267098_label.txt",
            )
            output = root / "out"
            output.mkdir()
            runner.write_public_result_txt(records, output / "result.txt")
            self.assertEqual(
                (output / "result.txt").read_text(encoding="utf-8"),
                "23_asr_1782715267098_1118_2308 speaker_02 \u6765\u5427\n"
                "23_asr_1782715267098_3426_6846 speaker_00 \u53ef\u4ee5\n",
            )
            runner.write_internal_asr_files(records, output / "inputs")
            self.assertTrue((output / "inputs" / "23_asr_1782715267098_asr.txt").is_file())

    def test_explicit_label_requires_a_single_result_audio(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            records = [
                runner.ResultRecord(Path("one/result.json"), "one.pcm", "one", []),
                runner.ResultRecord(Path("two/result.json"), "two.pcm", "two", []),
            ]
            explicit = root / "one_label.txt"
            explicit.write_text("one_0_1000 (speaker) text\n", encoding="utf-8")
            with self.assertRaisesRegex(ValueError, "--label"):
                runner.resolve_labels(records, None, explicit)


class SegmentMappingTest(unittest.TestCase):
    def _segment(
        self,
        segment_id: str,
        speaker_id: str,
        text: str,
    ) -> runner.TimedSegment:
        file_id, start_ms, duration_ms = runner.parse_segment_id(segment_id)
        return runner.TimedSegment(file_id, segment_id, start_ms, duration_ms, speaker_id, text)

    def test_whole_audio_wer_inputs_and_overlap_aware_mapping(self) -> None:
        predictions = [
            self._segment("audio_0_1000", "speaker_00", "\u7532"),
            self._segment("audio_1000_1000", "speaker_01", "\u4e59\u4e19"),
            self._segment("audio_3000_1000", "speaker_02", "\u4e01"),
        ]
        references = [
            self._segment("audio_0_500", "alice", "\u7532"),
            self._segment("audio_500_500", "multi", "\u91cd\u53e0"),
            self._segment("audio_1000_500", "bob", "\u4e59"),
            self._segment("audio_1500_500", "bob", "\u4e19"),
            self._segment("audio_2200_300", "alice", "\u620a"),
        ]
        record = runner.ResultRecord(Path("run/result.json"), "audio.pcm", "audio", predictions)
        labels = {"audio": runner.LabelRecord(Path("audio_label.txt"), "audio", references)}
        label_lines, hyp_lines = runner.build_whole_audio_wer_inputs([record], labels)
        self.assertEqual(label_lines, ["audio \u7532\u91cd\u53e0\u4e59\u4e19\u620a"])
        self.assertEqual(hyp_lines, ["audio \u7532\u4e59\u4e19\u4e01"])

        rows = runner.build_segment_detail_rows(references, predictions, "ZH")
        self.assertEqual(rows[0]["match_status"], "matched")
        self.assertEqual(rows[0]["mapping_type"], "one_to_one")
        self.assertEqual(rows[1]["match_status"], "overlap_not_scored")
        self.assertIsNone(rows[1]["segment_wer_percent"])
        self.assertEqual(rows[2]["mapping_type"], "one_to_many")
        self.assertEqual(rows[3]["mapping_type"], "one_to_many")
        self.assertEqual(rows[4]["match_status"], "unmatched_reference")
        self.assertEqual(rows[-1]["match_status"], "unmatched_prediction")


class RunnerIntegrationTest(unittest.TestCase):
    def test_cli_writes_standardized_evaluation_artifacts(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            run = root / "run"
            labels = root / "labels"
            output = root / "evaluation"
            run.mkdir()
            labels.mkdir()
            (run / "result.json").write_text(
                json.dumps(
                    {
                        "audio_name": "meeting.pcm",
                        "segments": [
                            {
                                "segment_id": "meeting_0_1000",
                                "duration_ms": 1000,
                                "speaker_id": "speaker_00",
                                "asr_text": "\u4f60\u597d",
                            },
                            {
                                "segment_id": "meeting_1000_1000",
                                "duration_ms": 1000,
                                "speaker_id": "speaker_01",
                                "asr_text": "\u4e16\u754c",
                            },
                        ],
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            (labels / "meeting_label.txt").write_text(
                "meeting_0_1000 (alice) \u4f60\u597d\n"
                "meeting_1000_1000 (bob) \u4e16\u754c\n",
                encoding="utf-8",
            )
            exit_code = runner.main(
                [
                    "--results-dir", str(run),
                    "--labels-dir", str(labels),
                    "--output-dir", str(output),
                    "--language", "ZH",
                    "--boundary-tolerance-ms", "500",
                    "--collar-ms", "0",
                ]
            )
            self.assertEqual(exit_code, 0)
            self.assertTrue((output / "result.txt").is_file())
            self.assertTrue((output / "asr" / "wer_detail.txt").is_file())
            self.assertTrue((output / "asr" / "wer_summary.json").is_file())
            self.assertTrue((output / "speaker" / "speaker_diarization_summary.json").is_file())
            self.assertTrue((output / "speaker" / "speaker_diarization_boundary_details.csv").is_file())
            self.assertTrue((output / "asr_segment_diff.xlsx").is_file())
            self.assertTrue((output / "evaluation_manifest.json").is_file())
            self.assertTrue((output / "evaluation_report.md").is_file())
            status = json.loads((output / "evaluation_status.json").read_text(encoding="utf-8"))
            self.assertEqual(status["tasks"]["wer"]["status"], "success")
            self.assertEqual(status["tasks"]["speaker"]["status"], "success")
            self.assertEqual(status["tasks"]["excel"]["status"], "success")
            from openpyxl import load_workbook

            workbook = load_workbook(output / "asr_segment_diff.xlsx", read_only=True)
            try:
                self.assertEqual(workbook.sheetnames, ["summary", "segment_details"])
            finally:
                workbook.close()

    def test_missing_label_writes_normalized_result_and_skips_metrics(self) -> None:
        with tempfile.TemporaryDirectory() as temp_dir:
            root = Path(temp_dir)
            run = root / "run"
            labels = root / "labels"
            output = root / "evaluation"
            run.mkdir()
            labels.mkdir()
            (run / "result.json").write_text(
                json.dumps(
                    {
                        "audio_name": "unlabeled.pcm",
                        "segments": [
                            {
                                "segment_id": "unlabeled_0_1000",
                                "duration_ms": 1000,
                                "speaker_id": "speaker_00",
                                "asr_text": "text",
                            }
                        ],
                    }
                ),
                encoding="utf-8",
            )
            self.assertEqual(
                runner.main(
                    [
                        "--results-dir", str(run),
                        "--labels-dir", str(labels),
                        "--output-dir", str(output),
                    ]
                ),
                0,
            )
            status = json.loads((output / "evaluation_status.json").read_text(encoding="utf-8"))
            self.assertEqual(status["tasks"]["wer"]["status"], "skipped")
            self.assertEqual(status["tasks"]["speaker"]["status"], "skipped")
            self.assertTrue((output / "result.txt").is_file())


if __name__ == "__main__":
    unittest.main()
